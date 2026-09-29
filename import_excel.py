"""Importa in Supabase i movimenti dai file Excel.

Uso:
    python3 import_excel.py "Portafogli.xlsx" "Storico/Portafogli 2025.xlsx" ...
    python3 import_excel.py --dry-run "Portafogli.xlsx"     # legge e riassume, non scrive

Cosa legge da ogni file:
  - fogli "Uscite" ed "Entrate" (Data, Operazione, Dettagli, Categoria, Importo, Note)
    oppure, se mancano, il primo foglio con una riga di intestazione che contiene
    "Data" e "Importo" (es. estratto conto esportato dalla banca);
  - "Dashboard Riassuntiva": budget mensile per categoria e saldo del conto a inizio anno;
  - "Investimenti Dashboard": registro delle operazioni sugli investimenti (le righe con
    data =OGGI()+1 diventano operazioni "domani", sempre con la data di domani);
  - "Pre 2023" (Anno | Mesi lavorati | Guadagno | Spesa | ... | Conto corrente): riepiloghi
    annuali e saldi a inizio anno (questi ultimi solo se mancano nel database).

Si può rilanciare sugli stessi file: i movimenti già presenti non vengono duplicati
(stessa data, importo, operazione e dettagli), così come le operazioni sugli
investimenti già presenti; budget e saldo vengono aggiornati.
Le categorie della banca vengono convertite con la mappatura salvata nel database
(tabella mappatura_categorie, modificabile dal sito in Movimenti).
URL e chiave vengono letti da assets/config.js.
"""
import argparse
import datetime as dt
import json
import re
import urllib.error
import urllib.parse
import urllib.request
import warnings
from collections import Counter
from pathlib import Path

import openpyxl

CONFIG_FILE = Path(__file__).parent / "assets" / "config.js"


# ---------------------------------------------------------------------------
# Supabase (REST)
# ---------------------------------------------------------------------------

def read_config():
    text = CONFIG_FILE.read_text(encoding="utf-8")
    url = re.search(r'supabaseUrl:\s*"([^"]*)"', text).group(1)
    key = re.search(r'supabaseKey:\s*"([^"]*)"', text).group(1)
    return url.rstrip("/"), key


class Supabase:
    def __init__(self, url, key):
        self.url = url
        self.key = key

    def request(self, method, path, body=None, prefer=None):
        request = urllib.request.Request(
            f"{self.url}/rest/v1/{path}",
            data=json.dumps(body).encode("utf-8") if body is not None else None,
            method=method,
        )
        request.add_header("apikey", self.key)
        request.add_header("Content-Type", "application/json")
        if prefer:
            request.add_header("Prefer", prefer)
        try:
            with urllib.request.urlopen(request) as response:
                raw = response.read()
        except urllib.error.HTTPError as error:
            raise SystemExit(f"Errore Supabase {error.code} su {path}: {error.read().decode()}") from None
        return json.loads(raw) if raw else None

    def select_all(self, table, columns="*"):
        rows, offset = [], 0
        while True:
            page = self.request("GET", f"{table}?select={columns}&order=id&limit=1000&offset={offset}")
            rows += page
            if len(page) < 1000:
                return rows
            offset += 1000

    def insert(self, table, rows):
        created = []
        for start in range(0, len(rows), 500):
            created += self.request("POST", table, rows[start:start + 500], prefer="return=representation")
        return created

    def upsert(self, table, rows, on_conflict, ignora_esistenti=False):
        """Inserisce o aggiorna; con ignora_esistenti=True le righe già presenti restano com'erano."""
        if rows:
            self.request(
                "POST",
                f"{table}?on_conflict={urllib.parse.quote(on_conflict)}",
                rows,
                prefer=f"resolution={'ignore' if ignora_esistenti else 'merge'}-duplicates,return=minimal",
            )


# ---------------------------------------------------------------------------
# Lettura Excel
# ---------------------------------------------------------------------------

def text(value):
    return "" if value is None else str(value).strip()


def to_date(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    match = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", text(value))
    if match:
        day, month, year = map(int, match.groups())
        return dt.date(year, month, day).isoformat()
    return None


def to_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    value = text(value).replace("€", "").replace(" ", "")
    if "," in value:
        value = value.replace(".", "").replace(",", ".")
    try:
        return float(value)
    except ValueError:
        return None


def read_table(sheet, header_row, default_tipo=None):
    """Legge i movimenti sotto la riga di intestazione (nomi colonna senza spazi finali)."""
    header = [text(cell.value).lower() for cell in sheet[header_row]]
    column = {name: index for index, name in enumerate(header) if name}
    movimenti = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        get = lambda name: row[column[name]] if name in column and column[name] < len(row) else None  # noqa: E731
        data, importo = to_date(get("data")), to_number(get("importo"))
        if not data or importo is None or importo == 0:
            continue
        categoria = text(get("categoria"))
        movimenti.append({
            "data": data,
            "operazione": text(get("operazione")),
            "dettagli": text(get("dettagli")),
            "categoria": categoria,
            "tipo": default_tipo or ("uscita" if importo < 0 else "entrata"),
            "importo": round(importo, 2),
            "note": text(get("note")) or None,
        })
    return movimenti


def read_movimenti(workbook):
    if "Uscite" in workbook.sheetnames or "Entrate" in workbook.sheetnames:
        movimenti = []
        for name, tipo in (("Uscite", "uscita"), ("Entrate", "entrata")):
            if name in workbook.sheetnames:
                movimenti += read_table(workbook[name], 1, tipo)
        return movimenti
    # Estratto conto: cerca la riga di intestazione nelle prime 40 righe
    sheet = workbook.worksheets[0]
    for row_number, row in enumerate(sheet.iter_rows(max_row=40, values_only=True), start=1):
        names = {text(value).lower() for value in row}
        if "data" in names and "importo" in names:
            return read_table(sheet, row_number)
    return []


def read_dashboard(workbook):
    """Budget mensili (> 0) e saldo iniziale dal foglio "Dashboard Riassuntiva"."""
    if "Dashboard Riassuntiva" not in workbook.sheetnames:
        return {}, None
    sheet = workbook["Dashboard Riassuntiva"]
    budget = {}
    if text(sheet["U1"].value) == "Categoria":
        # Dal 2024: tabella Categoria/Budget nelle colonne U:V
        for row in range(2, sheet.max_row + 1):
            nome, valore = text(sheet.cell(row, 21).value), to_number(sheet.cell(row, 22).value)
            if nome and valore:
                budget[nome] = valore
    elif text(sheet["Q1"].value).upper() == "BUDGET":
        # 2023: categoria in colonna A, budget in colonna Q
        for row in range(2, sheet.max_row + 1):
            nome, valore = text(sheet.cell(row, 1).value), to_number(sheet.cell(row, 17).value)
            if nome and valore and not nome.startswith("Totale"):
                budget[nome] = valore

    saldo = None
    for row in range(1, 5):
        etichetta = text(sheet.cell(row, 1).value).lower()
        valore = to_number(sheet.cell(row + 1, 1).value)
        if etichetta.startswith("saldo conto") and "+" not in etichetta and valore is not None:
            saldo = round(valore, 2)
    return budget, saldo


def read_investimenti(workbook, formule):
    """Registro del foglio "Investimenti Dashboard" (intestazione ID | Data | Nome | ...).
    Le righe con data =OGGI()+1 (formule: lo stesso file letto con le formule) diventano
    operazioni "domani": per il sito la loro data è sempre domani."""
    if "Investimenti Dashboard" not in workbook.sheetnames:
        return []
    sheet = workbook["Investimenti Dashboard"]
    date_formule = [row[0] for row in formule["Investimenti Dashboard"].iter_rows(min_col=2, max_col=2, values_only=True)]
    domani = (dt.date.today() + dt.timedelta(days=1)).isoformat()
    righe, header_trovato = [], False
    for numero, row in enumerate(sheet.iter_rows(max_col=18, values_only=True)):
        if not header_trovato:
            header_trovato = text(row[0]) == "ID" and text(row[1]) == "Data" and text(row[2]) == "Nome"
            continue
        data, importo = to_date(row[1]), to_number(row[8])
        if row[0] is None or not data or importo is None:
            continue
        is_domani = "TODAY()" in text(date_formule[numero]).upper()
        righe.append({
            "posizione": int(row[0]),
            "data": domani if is_domani else data,
            "domani": is_domani,
            "nome": text(row[2]),
            "isin": text(row[3]) if text(row[3]) not in ("", "-") else None,
            "prodotto": text(row[4]) or None,
            "tipo": text(row[5]) or None,
            "operazione": text(row[6]),
            "quantita": to_number(row[7]),
            "importo": round(importo, 4),
            "commissioni": to_number(row[9]),
            "tassa": to_number(row[10]),
            "note": text(row[17]) or None,
        })
    return righe


def read_riepiloghi(workbook):
    """Foglio "Pre 2023": Anno | Mesi lavorati | Guadagno annuo totale | Spesa totale | ... | Conto corrente."""
    for sheet in workbook.worksheets:
        if text(sheet["A1"].value) != "Anno" or text(sheet["B1"].value) != "Mesi lavorati":
            continue
        header = [text(cell.value).lower() for cell in sheet[1]]
        col = {nome: header.index(nome) for nome in ("guadagno annuo totale", "spesa totale", "conto corrente") if nome in header}
        riepiloghi = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            anno = to_number(row[0])
            if not anno:
                continue
            riepiloghi.append({
                "anno": int(anno),
                "mesi_lavorati": to_number(row[1]),
                "entrate": to_number(row[col["guadagno annuo totale"]]),
                "uscite": to_number(row[col["spesa totale"]]),
                "conto": to_number(row[col["conto corrente"]]) if "conto corrente" in col else None,
            })
        return riepiloghi
    return []


def saldi_da_riepiloghi(riepiloghi):
    """Saldo a inizio anno: il conto a fine anno precedente; per il primo anno, conto - risparmio."""
    saldi = {}
    for r in sorted(riepiloghi, key=lambda r: r["anno"]):
        if r["conto"] is None:
            continue
        saldi[r["anno"] + 1] = round(r["conto"], 2)
        if r["anno"] not in saldi and r["entrate"] is not None and r["uscite"] is not None:
            saldi[r["anno"]] = round(r["conto"] - (r["entrate"] - r["uscite"]), 2)
    return saldi


def workbook_year(movimenti):
    years = Counter(m["data"][:4] for m in movimenti)
    return int(years.most_common(1)[0][0]) if years else None


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def movimento_key(data, importo, operazione, dettagli):
    return (data, f"{float(importo):.2f}", operazione.strip().lower(), dettagli.strip().lower())


def investimento_key(posizione, data, operazione, importo):
    return (int(posizione), data, operazione, f"{float(importo):.2f}")


def main():
    parser = argparse.ArgumentParser(description="Importa i file Excel in MyFinance (Supabase).")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--dry-run", action="store_true", help="legge i file e riassume senza scrivere")
    args = parser.parse_args()
    warnings.simplefilter("ignore")

    files = []
    riepiloghi = []
    for path in args.files:
        workbook = openpyxl.load_workbook(path, data_only=True)
        riepiloghi_file = read_riepiloghi(workbook)
        if riepiloghi_file:
            riepiloghi += riepiloghi_file
            print(f"{path.name}: riepiloghi annuali {', '.join(str(r['anno']) for r in riepiloghi_file)}, "
                  f"saldi a inizio anno {saldi_da_riepiloghi(riepiloghi_file)}")
            continue
        movimenti = read_movimenti(workbook)
        year = workbook_year(movimenti)
        budget, saldo = read_dashboard(workbook)
        investimenti = read_investimenti(workbook, openpyxl.load_workbook(path) if "Investimenti Dashboard" in workbook.sheetnames else None)
        files.append((path, movimenti, year, budget, saldo, investimenti))
        entrate = sum(m["importo"] for m in movimenti if m["tipo"] == "entrata")
        uscite = -sum(m["importo"] for m in movimenti if m["tipo"] == "uscita")
        print(f"{path.name}: anno {year}, {len(movimenti)} movimenti "
              f"(entrate {entrate:.2f}, uscite {uscite:.2f}), {len(budget)} budget, saldo iniziale {saldo}, "
              f"{len(investimenti)} operazioni sugli investimenti")

    if args.dry_run:
        return

    url, key = read_config()
    if not url or not key:
        raise SystemExit("Imposta supabaseUrl e supabaseKey in assets/config.js.")
    db = Supabase(url, key)

    # Categorie: la mappatura salvata nel database (tabella mappatura_categorie,
    # modificabile dal sito) converte le categorie della banca nelle proprie;
    # le categorie che restano sconosciute vengono create.
    categorie = {(c["tipo"], c["nome"]): c["id"] for c in db.select_all("categorie")}
    mappatura = {r["categoria_banca"]: r["categoria_id"] for r in db.select_all("mappatura_categorie")}
    nuove = {(m["tipo"], m["categoria"]) for _, movimenti, *_ in files for m in movimenti
             if m["categoria"] and m["categoria"] not in mappatura}
    nuove |= {("uscita", nome) for _, _, _, budget, _, _ in files for nome in budget if nome not in mappatura}
    nuove = [{"tipo": tipo, "nome": nome} for tipo, nome in sorted(nuove) if (tipo, nome) not in categorie]
    for categoria in db.insert("categorie", nuove):
        categorie[(categoria["tipo"], categoria["nome"])] = categoria["id"]
    print(f"Categorie create: {len(nuove)}" + (f" ({', '.join(c['nome'] for c in nuove)})" if nuove else ""))

    # Movimenti: inserisce solo quelli non ancora presenti
    presenti = Counter(
        movimento_key(r["data"], r["importo"], r["operazione"], r["dettagli"])
        for r in db.select_all("movimenti", "data,importo,operazione,dettagli")
    )
    # Le operazioni "domani" si confrontano senza la data (che cambia ogni giorno) e senza
    # l'importo (lo aggiorna la Edge Function "aggiorna-quotazioni")
    def chiave_investimento(r):
        if r["domani"]:
            return (int(r["posizione"]), "domani", r["operazione"])
        return investimento_key(r["posizione"], r["data"], r["operazione"], r["importo"])

    investimenti_presenti = Counter(
        chiave_investimento(r) for r in db.select_all("investimenti", "posizione,data,domani,operazione,importo")
    )
    # Ogni file si confronta con il database com'è in quel momento (movimenti già presenti
    # più quelli inseriti dai file precedenti): due file che si sovrappongono non creano doppioni
    for path, movimenti, year, budget, saldo, investimenti in files:
        disponibili = presenti.copy()
        da_inserire = []
        for m in movimenti:
            key = movimento_key(m["data"], m["importo"], m["operazione"], m["dettagli"])
            if disponibili[key] > 0:
                disponibili[key] -= 1
                continue
            da_inserire.append({
                "data": m["data"],
                "operazione": m["operazione"],
                "dettagli": m["dettagli"],
                "categoria_id": mappatura.get(m["categoria"]) or categorie.get((m["tipo"], m["categoria"])),
                "importo": m["importo"],
                "note": m["note"],
            })
        db.insert("movimenti", da_inserire)
        presenti.update(movimento_key(r["data"], r["importo"], r["operazione"], r["dettagli"]) for r in da_inserire)

        db.upsert("budget", [
            {"categoria_id": mappatura.get(nome) or categorie[("uscita", nome)], "anno": year, "importo_mensile": valore}
            for nome, valore in budget.items()
        ], "categoria_id,anno")
        if saldo is not None and year:
            db.upsert("saldi", [{"anno": year, "saldo_iniziale": saldo}], "anno")

        investimenti_disponibili = investimenti_presenti.copy()
        nuovi_investimenti = []
        for riga in investimenti:
            key = chiave_investimento(riga)
            if investimenti_disponibili[key] > 0:
                investimenti_disponibili[key] -= 1
            else:
                nuovi_investimenti.append(riga)
        db.insert("investimenti", nuovi_investimenti)
        investimenti_presenti.update(chiave_investimento(r) for r in nuovi_investimenti)

        print(f"{path.name}: {len(da_inserire)} movimenti nuovi, "
              f"{len(movimenti) - len(da_inserire)} già presenti, budget {year} aggiornati: {len(budget)}, "
              f"{len(nuovi_investimenti)} operazioni investimenti nuove")

    # Riepiloghi annuali (Pre 2023) e saldi a inizio anno che ne derivano:
    # i saldi già presenti nel database non vengono sovrascritti
    if riepiloghi:
        db.upsert("riepiloghi_annuali", [
            {"anno": r["anno"], "mesi_lavorati": r["mesi_lavorati"], "entrate": r["entrate"], "uscite": r["uscite"]}
            for r in riepiloghi
        ], "anno")
        saldi = saldi_da_riepiloghi(riepiloghi)
        db.upsert("saldi", [{"anno": anno, "saldo_iniziale": saldo} for anno, saldo in saldi.items()],
                  "anno", ignora_esistenti=True)
        print(f"Riepiloghi annuali salvati: {len(riepiloghi)}; saldi a inizio anno (se mancavano): {saldi}")


if __name__ == "__main__":
    main()
